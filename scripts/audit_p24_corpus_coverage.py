"""P24-A: trace known compound evidence gaps from raw source to P22 candidates.

The audit intentionally performs no retrieval, prompt, or HCX call.  It only
checks whether an already-known required fact is available at each stage of
the existing pipeline.
"""

from __future__ import annotations

import argparse
import json
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.evaluation.corpus_coverage_audit import coverage_owner, owner_counts


@dataclass(frozen=True)
class AuditCase:
    question_id: str
    source_path: str
    source_format: str
    raw_pages: tuple[int, ...]
    raw_terms: tuple[str, ...]
    corpus_chunk_ids: tuple[str, ...]
    semantic_candidate_ids: tuple[str, ...]
    candidate_evidence_complete: bool
    selected_support_ids: tuple[str, ...]
    note: str


CASES = (
    AuditCase(
        question_id="R-010",
        source_path="docs_renamed/doc20.docx",
        source_format="docx",
        raw_pages=(),
        raw_terms=("연금외수령", "기타소득세", "부득이한 사유"),
        corpus_chunk_ids=(
            "e8d7e6a69504e042-paragraph_group-f8ee1dd5d6ec",
            "e8d7e6a69504e042-table-05de077f1213",
        ),
        semantic_candidate_ids=("37adcdadd2f3f902-table-08a7c9924f4f",),
        candidate_evidence_complete=True,
        selected_support_ids=("37adcdadd2f3f902-table-08a7c9924f4f",),
        note="The selected table expresses the exception as '부득이한 연금외수령 사유', while the slot required the narrower phrase '해지' + '부득이한 사유'.",
    ),
    AuditCase(
        question_id="R-024",
        source_path="투자설명서/KR5144420081/R2_KR5144420081.pdf",
        source_format="pdf",
        raw_pages=(1, 5),
        raw_terms=("4등급", "보통 위험", "미국 달러화로 표시되는 채권"),
        corpus_chunk_ids=(
            "42226bacb5f3829a-paragraph_group-7dbc09f70d2b",
            "42226bacb5f3829a-table-5d852ad4c548",
        ),
        semantic_candidate_ids=("42226bacb5f3829a-table-9748bbdce95b",),
        candidate_evidence_complete=False,
        selected_support_ids=("42226bacb5f3829a-table-9748bbdce95b",),
        note="Investment target was selected, but the direct risk-grade chunk was absent from the P22 candidate set.",
    ),
    AuditCase(
        question_id="R-028",
        source_path="투자설명서/KR510902511M/R2_KR510902511M.pdf",
        source_format="pdf",
        raw_pages=(17,),
        raw_terms=("장기성장주", "국내 주식", "운용 전략"),
        corpus_chunk_ids=("ac97d050d2b39ec2-paragraph_group-b89270097331",),
        semantic_candidate_ids=("ac97d050d2b39ec2-table-fdab853f4031",),
        candidate_evidence_complete=True,
        selected_support_ids=("ac97d050d2b39ec2-table-fdab853f4031",),
        note="A selected semantic-equivalent chunk contains both the domestic-stock target and strategy, but it lacks the literal '투자대상' field label required by the matcher.",
    ),
    AuditCase(
        question_id="R-037",
        source_path="docs_renamed/doc11.pdf",
        source_format="pdf",
        raw_pages=(1,),
        raw_terms=("DB제도", "DC제도", "적립금 운용 주체", "근로자"),
        corpus_chunk_ids=("4607500e74afcdf8-table-6f164502cbce",),
        semantic_candidate_ids=(),
        candidate_evidence_complete=False,
        selected_support_ids=(),
        note="The DB/DC operation comparison table exists in source and corpus but no semantic equivalent entered the P22 candidate set.",
    ),
)


def _nfd(value: str) -> str:
    return unicodedata.normalize("NFD", value)


def resolve_raw_source(raw_root: Path, source_path: str) -> Path:
    """Resolve paths despite macOS decomposed Korean filenames."""

    expected = _nfd(source_path)
    for path in raw_root.rglob("*"):
        if path.is_file() and _nfd(path.relative_to(raw_root).as_posix()) == expected:
            return path
    raise FileNotFoundError(f"Raw source not found: {source_path}")


def _docx_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_rows = [
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + table_rows)


def _pdf_text(path: Path, pages: tuple[int, ...]) -> str:
    document = fitz.open(path)
    try:
        return "\n".join(document[page - 1].get_text("text") for page in pages)
    finally:
        document.close()


def _raw_text(case: AuditCase, path: Path) -> str:
    if case.source_format == "docx":
        return _docx_text(path)
    if case.source_format == "pdf":
        return _pdf_text(path, case.raw_pages)
    raise ValueError(f"Unsupported audit source format: {case.source_format}")


def _load_chunks(path: Path) -> dict[str, dict]:
    return {
        chunk["chunk_id"]: chunk
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
        for chunk in (json.loads(line),)
    }


def _load_p22_rows(path: Path) -> dict[str, dict]:
    run = json.loads(path.read_text(encoding="utf-8"))
    return {row["question_id"]: row for row in run["rows"]}


def audit_case(case: AuditCase, *, raw_root: Path, chunks: dict[str, dict], p22_rows: dict[str, dict]) -> dict:
    raw_path = resolve_raw_source(raw_root, case.source_path)
    raw_text = _raw_text(case, raw_path)
    raw_term_found = {term: term in raw_text for term in case.raw_terms}
    corpus_found = {chunk_id: chunk_id in chunks for chunk_id in case.corpus_chunk_ids}
    row = p22_rows[case.question_id]
    candidate_ids = row["candidate_chunk_ids"]
    selected_ids = row["selected_merged_evidence_ids"]
    semantic_candidate_found = {
        chunk_id: chunk_id in candidate_ids for chunk_id in case.semantic_candidate_ids
    }
    selected_support_found = {
        chunk_id: chunk_id in selected_ids for chunk_id in case.selected_support_ids
    }
    raw_evidence_exists = all(raw_term_found.values())
    corpus_evidence_exists = all(corpus_found.values())
    semantic_evidence_in_candidates = (
        case.candidate_evidence_complete
        and bool(case.semantic_candidate_ids)
        and all(semantic_candidate_found.values())
    )
    selected_evidence_supports_requirement = bool(case.selected_support_ids) and all(selected_support_found.values())
    owner = coverage_owner(
        raw_evidence_exists=raw_evidence_exists,
        corpus_evidence_exists=corpus_evidence_exists,
        semantic_evidence_in_candidates=semantic_evidence_in_candidates,
        selected_evidence_supports_requirement=selected_evidence_supports_requirement,
    )
    # A selected semantic candidate proves recall for R-010/R-028.  The final
    # gate still failed because it could not match the candidate to its slot.
    if owner == "selection":
        owner = "evidence_matcher"
    return {
        "question_id": case.question_id,
        "source": {
            "relative_path": case.source_path,
            "format": case.source_format,
            "raw_file_found": raw_path.exists(),
            "raw_pages_checked": list(case.raw_pages),
            "native_text_terms_found": raw_term_found,
        },
        "corpus": {
            "required_chunk_ids": list(case.corpus_chunk_ids),
            "required_chunks_found": corpus_found,
        },
        "p22_retrieval": {
            "candidate_chunk_ids": candidate_ids,
            "selected_merged_evidence_ids": selected_ids,
            "semantic_candidate_ids": list(case.semantic_candidate_ids),
            "semantic_candidates_in_p22": semantic_candidate_found,
            "semantic_candidates_cover_all_requirements": case.candidate_evidence_complete,
            "selected_support_ids": list(case.selected_support_ids),
            "selected_support_in_p22": selected_support_found,
            "missing_requirement_slots": row["missing_requirement_slots"],
        },
        "primary_owner": owner,
        "finding": case.note,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-root", type=Path, default=ROOT / "data/raw/연금")
    parser.add_argument("--chunks", type=Path, default=ROOT / "data/parsed/chunks.jsonl")
    parser.add_argument("--p22-run", type=Path, default=ROOT / "data/diagnostics/p22_full_hcx.json")
    parser.add_argument("--output", type=Path, default=ROOT / "evaluation/p24_corpus_coverage_audit.json")
    args = parser.parse_args()

    chunks = _load_chunks(args.chunks)
    p22_rows = _load_p22_rows(args.p22_run)
    records = [audit_case(case, raw_root=args.raw_root, chunks=chunks, p22_rows=p22_rows) for case in CASES]
    payload = {
        "audit": "P24-A corpus coverage audit",
        "scope": [case.question_id for case in CASES],
        "method": "raw source -> parsed corpus -> P22 candidate/selected evidence; no HCX call",
        "owner_counts": owner_counts(records),
        "records": records,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "owner_counts": payload["owner_counts"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
