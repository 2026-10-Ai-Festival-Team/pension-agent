"""Run lightweight, read-only structural diagnostics for representative documents."""

import csv
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Dict, Iterable, List

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from check_data_path import get_data_root


SAMPLES_PATH = Path("data/representative_documents.csv")
REQUIREMENTS_PATH = Path("data/representative_document_requirements.csv")
DIAGNOSTICS_ROOT = Path("data/diagnostics")
REPORT_PATH = Path("data/parse_quality_report.md")
DATE_PATTERN = re.compile(r"\b(?:20\d{2}[.\-/년\s]+\d{1,2}[.\-/월\s]+\d{1,2}일?|20\d{2}[.\-/]\d{1,2})\b")
PRODUCT_CODE_PATTERN = re.compile(r"\bKR[A-Z0-9]+\b")


def candidates(text: str, pattern: re.Pattern[str], limit: int = 10) -> List[str]:
    return list(dict.fromkeys(match.group(0) for match in pattern.finditer(text)))[:limit]


def first_nonempty_line(text: str) -> str:
    for line in text.splitlines():
        normalized = " ".join(line.split())
        if normalized:
            return normalized[:200]
    return ""


def classify_section(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    short_lines = sum(len(line) <= 80 for line in lines)
    return "heading" if lines and short_lines / len(lines) >= 0.25 else "flat"


def diagnose_pdf(path: Path) -> Dict[str, Any]:
    document = fitz.open(path)
    texts = [page.get_text("text") for page in document]
    pages_without_text = [index + 1 for index, text in enumerate(texts) if not text.strip()]
    image_count = sum(len(page.get_images(full=True)) for page in document)
    all_text = "\n".join(texts)
    if len(pages_without_text) == len(document):
        scan_likelihood = "full"
    elif pages_without_text:
        scan_likelihood = "partial"
    else:
        scan_likelihood = "none"
    table_candidates = 0
    for page in document:
        try:
            table_candidates += len(page.find_tables().tables)
        except (AttributeError, RuntimeError):
            pass
    result = {
        "page_count": len(document), "native_text_available": bool(all_text.strip()),
        "native_text_length": len(all_text), "pages_with_no_text": pages_without_text,
        "scan_likelihood": scan_likelihood, "has_tables": table_candidates > 0,
        "table_complexity": "manual_review" if table_candidates else "none_detected",
        "image_count": image_count, "has_images": image_count > 0,
        "title_candidates": [value for value in [document.metadata.get("title", ""), first_nonempty_line(all_text)] if value],
        "section_structure": classify_section(all_text), "page_number_preservable": True,
        "effective_date_candidates": candidates(all_text, DATE_PATTERN),
        "product_code_candidates": candidates(all_text, PRODUCT_CODE_PATTERN),
        "header_footer_noise": "manual_review", "recommended_parser": "pymupdf",
        "ocr_required": scan_likelihood != "none", "warnings": [],
    }
    if pages_without_text:
        result["warnings"].append(f"pages with no native text: {pages_without_text}")
    if table_candidates:
        result["warnings"].append(f"table candidates detected: {table_candidates}")
    document.close()
    return result


def diagnose_docx(path: Path) -> Dict[str, Any]:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    with zipfile.ZipFile(path) as archive:
        image_count = sum(name.startswith("word/media/") for name in archive.namelist())
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading")]
    return {
        "page_count": "not_available", "native_text_available": bool(text.strip()), "native_text_length": len(text),
        "scan_likelihood": "none", "has_tables": bool(document.tables),
        "table_complexity": "manual_review" if document.tables else "none_detected", "image_count": image_count,
        "has_images": image_count > 0, "title_candidates": [value for value in [document.core_properties.title, first_nonempty_line(text)] if value],
        "section_structure": "heading" if headings else classify_section(text), "page_number_preservable": False,
        "effective_date_candidates": candidates(text, DATE_PATTERN), "product_code_candidates": candidates(text, PRODUCT_CODE_PATTERN),
        "header_footer_noise": "manual_review", "recommended_parser": "python-docx", "ocr_required": False,
        "warnings": ["DOCX has no stable rendered page positions without a rendering step"],
    }


def diagnose_pptx(path: Path) -> Dict[str, Any]:
    presentation = Presentation(path)
    texts: List[str] = []
    table_count = image_count = 0
    title_candidates: List[str] = []
    for slide in presentation.slides:
        if slide.shapes.title and slide.shapes.title.text.strip():
            title_candidates.append(slide.shapes.title.text.strip())
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
            if getattr(shape, "has_table", False):
                table_count += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
    text = "\n".join(texts)
    return {
        "page_count": len(presentation.slides), "native_text_available": bool(text.strip()), "native_text_length": len(text),
        "scan_likelihood": "none", "has_tables": table_count > 0,
        "table_complexity": "manual_review" if table_count else "none_detected", "image_count": image_count,
        "has_images": image_count > 0, "title_candidates": list(dict.fromkeys(title_candidates))[:10] or [first_nonempty_line(text)],
        "section_structure": "slides", "page_number_preservable": True,
        "effective_date_candidates": candidates(text, DATE_PATTERN), "product_code_candidates": candidates(text, PRODUCT_CODE_PATTERN),
        "header_footer_noise": "manual_review", "recommended_parser": "python-pptx", "ocr_required": False,
        "warnings": ["Text-box reading order requires visual spot-checking"],
    }


def diagnose_xlsx(path: Path) -> Dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet_names = workbook.sheetnames
    text_values: List[str] = []
    table_count = image_count = merged_count = formula_count = 0
    hidden = []
    for worksheet in workbook.worksheets:
        merged_count += len(worksheet.merged_cells.ranges)
        table_count += len(worksheet.tables)
        image_count += len(worksheet._images)
        if worksheet.sheet_state != "visible":
            hidden.append(worksheet.title)
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    text_values.append(cell.value)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    text = "\n".join(text_values)
    warnings = []
    if hidden:
        warnings.append(f"hidden sheets: {hidden}")
    if merged_count:
        warnings.append(f"merged-cell ranges: {merged_count}")
    if formula_count:
        warnings.append(f"formula cells: {formula_count}")
    return {
        "page_count": f"sheets:{len(sheet_names)}", "native_text_available": bool(text.strip()), "native_text_length": len(text),
        "scan_likelihood": "none", "has_tables": table_count > 0,
        "table_complexity": "merged" if merged_count else ("simple" if table_count else "manual_review"),
        "image_count": image_count, "has_images": image_count > 0,
        "title_candidates": [workbook.properties.title] if workbook.properties.title else sheet_names[:1],
        "section_structure": "workbook", "page_number_preservable": "sheet_cell",
        "effective_date_candidates": candidates(text, DATE_PATTERN), "product_code_candidates": candidates(text, PRODUCT_CODE_PATTERN),
        "header_footer_noise": "not_applicable", "recommended_parser": "openpyxl", "ocr_required": False,
        "warnings": warnings,
    }


def diagnose(path: Path) -> Dict[str, Any]:
    dispatch = {".pdf": diagnose_pdf, ".docx": diagnose_docx, ".pptx": diagnose_pptx, ".xlsx": diagnose_xlsx}
    return dispatch[path.suffix.lower()](path)


def write_requirements(rows: Iterable[Dict[str, str]], diagnostics: Dict[str, Dict[str, Any]]) -> None:
    with REQUIREMENTS_PATH.open("w", encoding="utf-8-sig", newline="") as file:
        fieldnames = [
            "file_id", "relative_path", "filename", "sample_role", "file_format", "page_count", "native_text_available", "scan_likelihood", "has_tables", "table_complexity", "has_images", "title_extractable", "section_structure", "page_number_preservable", "effective_date_found", "product_code_found", "header_footer_noise", "recommended_parser", "ocr_required", "review_notes", "review_status"
        ]
        writer = csv.DictWriter(file, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            info = diagnostics[row["file_id"]]
            writer.writerow({
                "file_id": row["file_id"], "relative_path": row["relative_path"], "filename": row["filename"], "sample_role": row["sample_role"], "file_format": row["extension"],
                "page_count": info["page_count"], "native_text_available": info["native_text_available"], "scan_likelihood": info["scan_likelihood"], "has_tables": info["has_tables"], "table_complexity": info["table_complexity"], "has_images": info["has_images"], "title_extractable": bool(info["title_candidates"]), "section_structure": info["section_structure"], "page_number_preservable": info["page_number_preservable"], "effective_date_found": "; ".join(info["effective_date_candidates"]), "product_code_found": "; ".join(info["product_code_candidates"]) or row["product_code"], "header_footer_noise": info["header_footer_noise"], "recommended_parser": info["recommended_parser"], "ocr_required": info["ocr_required"], "review_notes": "; ".join(info["warnings"]), "review_status": "diagnosed",
            })


def write_report(rows: Iterable[Dict[str, str]], diagnostics: Dict[str, Dict[str, Any]]) -> None:
    rows = list(rows)
    parser_counts = Counter(diagnostics[row["file_id"]]["recommended_parser"] for row in rows)
    ocr_rows = [row for row in rows if diagnostics[row["file_id"]]["ocr_required"]]
    table_rows = [row for row in rows if diagnostics[row["file_id"]]["has_tables"]]
    lines = ["# Representative Parse Quality Report", "", "## Coverage", "", f"- Diagnosed representative documents: {len(rows)}", f"- OCR-required candidates: {len(ocr_rows)}", f"- Table-bearing candidates: {len(table_rows)}", "", "## Recommended parser trials", ""]
    lines.extend(f"- `{parser}`: {count}" for parser, count in sorted(parser_counts.items()))
    lines.extend(["", "## Documents needing extra attention", ""])
    for row in rows:
        info = diagnostics[row["file_id"]]
        if info["warnings"] or info["ocr_required"]:
            lines.append(f"- `{row['relative_path']}`: " + "; ".join(info["warnings"] or ["OCR required"]))
    if lines[-1] == "":
        lines.append("- None")
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    data_root = get_data_root()
    with SAMPLES_PATH.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    diagnostics: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        source_path = data_root / Path(row["relative_path"])
        info = diagnose(source_path)
        info.update({"file_id": row["file_id"], "relative_path": row["relative_path"], "extension": row["extension"]})
        diagnostics[row["file_id"]] = info
        output = DIAGNOSTICS_ROOT / row["extension"].lstrip(".") / f"{row['file_id']}.json"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(info, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_requirements(rows, diagnostics)
    write_report(rows, diagnostics)
    print(f"진단 완료: {len(rows)}개 문서")
    print(f"검토표: {REQUIREMENTS_PATH}")
    print(f"품질 보고서: {REPORT_PATH}")


if __name__ == "__main__":
    main()
