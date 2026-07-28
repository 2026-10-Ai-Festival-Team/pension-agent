"""DOCX parser that preserves body block order without inventing page locations."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.ingestion.common import (
    extract_date_candidates,
    extract_product_codes,
    make_source_id,
    normalize_text,
    relative_path,
)
from src.models.document import (
    DocumentElement,
    DocumentType,
    ElementType,
    Locator,
    ParsedDocument,
    SourceFormat,
    TableData,
)


class DocxParser:
    supported_extensions = {".docx"}

    def parse(self, path: Path, source_root: Path) -> ParsedDocument:
        relative = relative_path(path, source_root)
        source_id = make_source_id(relative)
        document = Document(path)
        elements: list[DocumentElement] = []
        full_text_parts: list[str] = []
        warnings = self._unsupported_content_warnings(document)

        for block_index, block in enumerate(document.iter_inner_content()):
            if isinstance(block, Paragraph):
                text = normalize_text(block.text)
                if not text:
                    continue
                elements.append(
                    DocumentElement(
                        element_id=f"{source_id}-b{block_index}",
                        order=len(elements),
                        kind=self._infer_paragraph_kind(block),
                        locator=Locator(block_index=block_index),
                        text=text,
                        metadata={
                            "style_name": block.style.name if block.style is not None else None,
                        },
                    )
                )
                full_text_parts.append(text)
            elif isinstance(block, Table):
                rows = [
                    [normalize_text(cell.text) or None for cell in row.cells]
                    for row in block.rows
                ]
                if self._has_duplicated_merged_cells(block):
                    warnings.append(
                        f"block {block_index}: merged table cells may be duplicated"
                    )
                elements.append(
                    DocumentElement(
                        element_id=f"{source_id}-t{block_index}",
                        order=len(elements),
                        kind=ElementType.TABLE,
                        locator=Locator(block_index=block_index),
                        table=TableData(rows=rows),
                        metadata={"parser": "python-docx"},
                    )
                )
                full_text_parts.extend(value for row in rows for value in row if value)

        full_text = "\n".join(full_text_parts)
        return ParsedDocument(
            source_id=source_id,
            relative_path=relative,
            filename=path.name,
            source_format=SourceFormat.DOCX,
            document_type=self._infer_document_type(relative),
            title=self._infer_title(elements),
            product_codes=extract_product_codes(relative, full_text),
            date_candidates=extract_date_candidates(full_text),
            elements=elements,
            warnings=warnings,
            metadata={"parser": "python-docx", "element_count": len(elements)},
        )

    @staticmethod
    def _infer_paragraph_kind(paragraph: Paragraph) -> ElementType:
        style_name = paragraph.style.name.lower() if paragraph.style is not None else ""
        if style_name in {"title", "제목"}:
            return ElementType.TITLE
        if "heading" in style_name or "제목" in style_name:
            return ElementType.HEADING
        return ElementType.PARAGRAPH

    @staticmethod
    def _infer_title(elements: list[DocumentElement]) -> str | None:
        for element in elements:
            if element.kind == ElementType.TITLE and element.text:
                return element.text
        for element in elements:
            if element.text and len(element.text) <= 150:
                return element.text.splitlines()[0]
        return None

    @staticmethod
    def _infer_document_type(relative: str) -> DocumentType:
        if "투자설명서" in relative:
            return DocumentType.INVESTMENT_PRODUCT
        if "docs_renamed" in relative:
            return DocumentType.PENSION_GUIDE
        return DocumentType.UNKNOWN

    @staticmethod
    def _has_duplicated_merged_cells(table: Table) -> bool:
        for row in table.rows:
            cell_ids = [id(cell._tc) for cell in row.cells]
            if len(cell_ids) != len(set(cell_ids)):
                return True
        return False

    @staticmethod
    def _unsupported_content_warnings(document: Document) -> list[str]:
        warnings: list[str] = []
        has_header_or_footer_text = any(
            normalize_text(paragraph.text)
            for section in document.sections
            for area in (section.header, section.footer)
            for paragraph in area.paragraphs
        )
        if has_header_or_footer_text:
            warnings.append("header/footer extraction not implemented")
        if document.element.body.xpath(".//w:txbxContent"):
            warnings.append("text box content may be omitted")
        return warnings
