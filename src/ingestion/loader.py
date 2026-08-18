from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".csv", ".json", ".xlsx", ".docx", ".pptx"}


@dataclass
class ParsedPage:
    page: int | None
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class LoadedDocument:
    path: Path
    document_id: str
    pages: list[ParsedPage]
    contains_table: bool = False


def discover_documents(raw_dir: Path) -> list[Path]:
    if not raw_dir.exists():
        return []
    return sorted(path for path in raw_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES)


def _document_id(path: Path, raw_dir: Path) -> str:
    import hashlib

    relative = path.relative_to(raw_dir).as_posix()
    return "DOC-" + hashlib.sha1(relative.encode("utf-8")).hexdigest()[:12].upper()


def load_document(path: Path, raw_dir: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    doc_id = _document_id(path, raw_dir)
    if suffix in {".txt", ".md"}:
        return LoadedDocument(path, doc_id, [ParsedPage(1, path.read_text("utf-8"))])
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF 처리를 위해 pypdf를 설치하세요.") from exc
        reader = PdfReader(str(path))
        pages = [ParsedPage(i, page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
        return LoadedDocument(path, doc_id, pages)
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        text = "\n".join(" | ".join(f"{key}: {value}" for key, value in row.items()) for row in rows)
        return LoadedDocument(path, doc_id, [ParsedPage(1, text)], contains_table=True)
    if suffix == ".json":
        data = json.loads(path.read_text("utf-8"))
        return LoadedDocument(path, doc_id, [ParsedPage(1, json.dumps(data, ensure_ascii=False, indent=2))])
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("XLSX 처리를 위해 openpyxl을 설치하세요.") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        pages: list[ParsedPage] = []
        for index, sheet in enumerate(workbook.worksheets, 1):
            lines = [" | ".join("" if cell is None else str(cell) for cell in row) for row in sheet.iter_rows(values_only=True)]
            pages.append(ParsedPage(index, "\n".join(lines), {"sheet_name": sheet.title}))
        return LoadedDocument(path, doc_id, pages, contains_table=True)
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX 처리를 위해 python-docx를 설치하세요.") from exc
        source = Document(path)
        blocks = [paragraph.text for paragraph in source.paragraphs if paragraph.text.strip()]
        for table in source.tables:
            blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return LoadedDocument(path, doc_id, [ParsedPage(1, "\n".join(blocks))], contains_table=bool(source.tables))
    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("PPTX 처리를 위해 python-pptx를 설치하세요.") from exc
        source = Presentation(path)
        pages = []
        for index, slide in enumerate(source.slides, 1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            pages.append(ParsedPage(index, "\n".join(texts)))
        return LoadedDocument(path, doc_id, pages)
    raise ValueError(f"지원하지 않는 파일 형식: {suffix}")
