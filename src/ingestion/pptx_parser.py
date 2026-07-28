"""PPTX parser preserving slide locations and recursively traversing groups."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.ingestion.common import (
    extract_date_candidates,
    extract_product_codes,
    make_source_id,
    normalize_text,
    relative_path,
)
from src.models.document import (
    DocumentElement,
    DocumentType,
    ElementType,
    Locator,
    ParsedDocument,
    SourceFormat,
    TableData,
)


class PptxParser:
    supported_extensions = {".pptx"}

    def parse(self, path: Path, source_root: Path) -> ParsedDocument:
        relative = relative_path(path, source_root)
        source_id = make_source_id(relative)
        presentation = Presentation(str(path))
        elements: List[DocumentElement] = []
        full_text_parts: List[str] = []
        warnings: List[str] = []
        image_count = chart_count = unsupported_graphic_count = 0

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_element_start = len(elements)
            title_shape_id = slide.shapes.title.shape_id if slide.shapes.title is not None else None
            flattened_shapes = list(self._iter_shapes(slide.shapes))
            flattened_shapes.sort(key=lambda item: self._shape_sort_key(item[1]))

            for block_index, (group_path, shape) in enumerate(flattened_shapes):
                bbox = self._extract_bbox(shape)
                metadata: Dict[str, Any] = {
                    "parser": "python-pptx",
                    "shape_name": getattr(shape, "name", None),
                    "shape_id": getattr(shape, "shape_id", None),
                    "shape_type": self._shape_type_name(shape),
                    "group_path": group_path,
                    "bbox_unit": "emu",
                    "reading_order": "top_left_heuristic",
                }

                if getattr(shape, "has_table", False):
                    rows = self._extract_table_rows(shape.table)
                    elements.append(
                        DocumentElement(
                            element_id=f"{source_id}-s{slide_number}-t{block_index}",
                            order=len(elements),
                            kind=ElementType.TABLE,
                            locator=Locator(slide=slide_number, block_index=block_index, bbox=bbox),
                            table=TableData(rows=rows),
                            metadata={
                                **metadata,
                                "row_count": len(rows),
                                "column_count": max((len(row) for row in rows), default=0),
                                "merged_cell_handling": "visual grid may contain duplicate or spanned cells",
                            },
                        )
                    )
                    full_text_parts.extend(cell for row in rows for cell in row if cell)
                    continue

                if getattr(shape, "has_chart", False):
                    chart_count += 1
                    warnings.append(self._warning(slide_number, shape, block_index, "chart content not extracted"))
                    continue

                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    warnings.append(self._warning(slide_number, shape, block_index, "image content not extracted"))
                    continue

                if getattr(shape, "has_text_frame", False):
                    text = normalize_text(shape.text)
                    if not text:
                        continue
                    kind = ElementType.TITLE if title_shape_id == shape.shape_id else ElementType.PARAGRAPH
                    elements.append(
                        DocumentElement(
                            element_id=f"{source_id}-s{slide_number}-b{block_index}",
                            order=len(elements),
                            kind=kind,
                            locator=Locator(slide=slide_number, block_index=block_index, bbox=bbox),
                            text=text,
                            metadata={**metadata, "paragraphs": self._extract_paragraph_metadata(shape)},
                        )
                    )
                    full_text_parts.append(text)
                    continue

                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GRAPHIC_FRAME:
                    unsupported_graphic_count += 1
                    warnings.append(self._warning(slide_number, shape, block_index, "unsupported graphic frame"))

            if len(elements) == slide_element_start:
                warnings.append(f"slide {slide_number}: extractable text or table not found")

        full_text = "\n".join(full_text_parts)
        return ParsedDocument(
            source_id=source_id,
            relative_path=relative,
            filename=path.name,
            source_format=SourceFormat.PPTX,
            document_type=self._infer_document_type(relative),
            title=self._infer_title(elements),
            product_codes=extract_product_codes(relative, full_text),
            date_candidates=extract_date_candidates(full_text),
            elements=elements,
            warnings=warnings,
            metadata={
                "parser": "python-pptx",
                "slide_count": len(presentation.slides),
                "element_count": len(elements),
                "image_count": image_count,
                "chart_count": chart_count,
                "unsupported_graphic_count": unsupported_graphic_count,
                "notes_extracted": False,
            },
        )

    def _iter_shapes(self, shapes: Iterable[Any], group_path: Tuple[str, ...] = ()) -> Iterable[Tuple[List[str], Any]]:
        for shape in sorted(list(shapes), key=self._shape_sort_key):
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                group_name = getattr(shape, "name", None) or f"group-{getattr(shape, 'shape_id', 'unknown')}"
                yield from self._iter_shapes(shape.shapes, group_path + (group_name,))
                continue
            yield list(group_path), shape

    @staticmethod
    def _shape_sort_key(shape: Any) -> Tuple[int, int, int]:
        return (
            int(getattr(shape, "top", 0) or 0),
            int(getattr(shape, "left", 0) or 0),
            int(getattr(shape, "shape_id", 0) or 0),
        )

    @staticmethod
    def _extract_bbox(shape: Any) -> Optional[Tuple[float, float, float, float]]:
        try:
            left = float(shape.left)
            top = float(shape.top)
            return left, top, float(shape.left + shape.width), float(shape.top + shape.height)
        except (AttributeError, TypeError, ValueError):
            return None

    @staticmethod
    def _extract_table_rows(table: Any) -> List[List[Optional[str]]]:
        return [[normalize_text(cell.text) or None for cell in row.cells] for row in table.rows]

    @staticmethod
    def _extract_paragraph_metadata(shape: Any) -> List[Dict[str, Any]]:
        return [
            {"paragraph_index": index, "text": text, "level": int(paragraph.level)}
            for index, paragraph in enumerate(shape.text_frame.paragraphs)
            for text in [normalize_text(paragraph.text)]
            if text
        ]

    @staticmethod
    def _shape_type_name(shape: Any) -> str:
        shape_type = getattr(shape, "shape_type", None)
        return "unknown" if shape_type is None else getattr(shape_type, "name", str(shape_type))

    @staticmethod
    def _warning(slide_number: int, shape: Any, block_index: int, message: str) -> str:
        return f"slide {slide_number}, shape {getattr(shape, 'name', block_index)}: {message}"

    @staticmethod
    def _infer_title(elements: List[DocumentElement]) -> Optional[str]:
        for element in elements:
            if element.kind == ElementType.TITLE and element.text:
                return element.text.splitlines()[0]
        for element in elements:
            if element.locator.slide == 1 and element.text and len(element.text) <= 150:
                return element.text.splitlines()[0]
        return None

    @staticmethod
    def _infer_document_type(relative: str) -> DocumentType:
        if "투자설명서" in relative:
            return DocumentType.INVESTMENT_PRODUCT
        if "docs_renamed" in relative:
            return DocumentType.PENSION_GUIDE
        return DocumentType.UNKNOWN
