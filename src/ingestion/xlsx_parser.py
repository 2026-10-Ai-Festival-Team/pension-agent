"""XLSX parser preserving worksheet locations, merged ranges, and formulas."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.ingestion.common import (
    extract_date_candidates,
    extract_product_codes,
    make_source_id,
    normalize_text,
    relative_path,
)
from src.models.document import (
    DocumentElement,
    DocumentType,
    ElementType,
    Locator,
    ParsedDocument,
    SourceFormat,
    TableData,
)


class XlsxParser:
    supported_extensions = {".xlsx"}

    def parse(self, path: Path, source_root: Path) -> ParsedDocument:
        relative = relative_path(path, source_root)
        source_id = make_source_id(relative)
        formula_workbook = load_workbook(str(path), read_only=False, data_only=False)
        value_workbook = load_workbook(str(path), read_only=False, data_only=True)
        elements: List[DocumentElement] = []
        warnings: List[str] = []
        full_text_parts: List[str] = []
        hidden_sheet_count = total_formula_count = 0

        try:
            for sheet_index, formula_sheet in enumerate(formula_workbook.worksheets, start=1):
                value_sheet = value_workbook[formula_sheet.title]
                if formula_sheet.sheet_state != "visible":
                    hidden_sheet_count += 1
                bounds = self._find_used_bounds(formula_sheet)
                if bounds is None:
                    warnings.append(f"sheet {formula_sheet.title}: non-empty cell not found")
                    continue
                min_row, min_column, max_row, max_column = bounds
                cell_range = self._make_cell_range(min_row, min_column, max_row, max_column)
                rows: List[List[Optional[str]]] = []
                formula_cells: Dict[str, Dict[str, Any]] = {}
                formatted_cells: Dict[str, Dict[str, Any]] = {}

                for row_index in range(min_row, max_row + 1):
                    parsed_row: List[Optional[str]] = []
                    for column_index in range(min_column, max_column + 1):
                        formula_cell = formula_sheet.cell(row=row_index, column=column_index)
                        value_cell = value_sheet.cell(row=row_index, column=column_index)
                        coordinate = formula_cell.coordinate
                        raw_value = formula_cell.value
                        cached_value = value_cell.value
                        is_formula = isinstance(raw_value, str) and raw_value.startswith("=")
                        if is_formula:
                            total_formula_count += 1
                            formula_cells[coordinate] = {
                                "formula": raw_value,
                                "cached_value": self._serialize_cell_value(cached_value),
                                "cached_value_available": cached_value is not None,
                            }
                            selected_value = cached_value if cached_value is not None else raw_value
                        else:
                            selected_value = raw_value
                        serialized_value = self._serialize_cell_value(selected_value)
                        parsed_row.append(serialized_value)
                        if serialized_value:
                            full_text_parts.append(serialized_value)
                        if formula_cell.number_format and formula_cell.number_format != "General":
                            formatted_cells[coordinate] = {
                                "number_format": formula_cell.number_format,
                                "data_type": formula_cell.data_type,
                                "is_date": bool(formula_cell.is_date),
                            }
                    rows.append(parsed_row)

                merged_ranges = sorted(str(item) for item in formula_sheet.merged_cells.ranges)
                hidden_rows = sorted(index for index, dimension in formula_sheet.row_dimensions.items() if dimension.hidden)
                hidden_columns = sorted(name for name, dimension in formula_sheet.column_dimensions.items() if dimension.hidden)
                chart_count = len(formula_sheet._charts)
                image_count = len(formula_sheet._images)
                if chart_count:
                    warnings.append(f"sheet {formula_sheet.title}: chart content not extracted ({chart_count})")
                if image_count:
                    warnings.append(f"sheet {formula_sheet.title}: image content not extracted ({image_count})")
                elements.append(
                    DocumentElement(
                        element_id=f"{source_id}-sheet{sheet_index}",
                        order=len(elements),
                        kind=ElementType.TABLE,
                        locator=Locator(sheet=formula_sheet.title, cell_range=cell_range, block_index=sheet_index - 1),
                        table=TableData(rows=rows, merged_ranges=merged_ranges),
                        metadata={
                            "parser": "openpyxl",
                            "sheet_index": sheet_index,
                            "sheet_state": formula_sheet.sheet_state,
                            "min_row": min_row,
                            "min_column": min_column,
                            "max_row": max_row,
                            "max_column": max_column,
                            "row_count": max_row - min_row + 1,
                            "column_count": max_column - min_column + 1,
                            "formula_cells": formula_cells,
                            "formatted_cells": formatted_cells,
                            "hidden_rows": hidden_rows,
                            "hidden_columns": hidden_columns,
                            "merged_range_count": len(merged_ranges),
                            "formula_count": len(formula_cells),
                            "chart_count": chart_count,
                            "image_count": image_count,
                        },
                    )
                )

            full_text = "\n".join(full_text_parts)
            return ParsedDocument(
                source_id=source_id,
                relative_path=relative,
                filename=path.name,
                source_format=SourceFormat.XLSX,
                document_type=self._infer_document_type(relative),
                title=self._infer_title(elements),
                product_codes=extract_product_codes(relative, full_text),
                date_candidates=extract_date_candidates(full_text),
                elements=elements,
                warnings=warnings,
                metadata={
                    "parser": "openpyxl",
                    "sheet_count": len(formula_workbook.worksheets),
                    "hidden_sheet_count": hidden_sheet_count,
                    "element_count": len(elements),
                    "formula_count": total_formula_count,
                    "formulas_calculated": False,
                    "charts_extracted": False,
                    "images_extracted": False,
                    "macros_extracted": False,
                },
            )
        finally:
            formula_workbook.close()
            value_workbook.close()

    @staticmethod
    def _find_used_bounds(worksheet: Any) -> Optional[Tuple[int, int, int, int]]:
        min_row = min_column = max_row = max_column = None
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                min_row = cell.row if min_row is None else min(min_row, cell.row)
                max_row = cell.row if max_row is None else max(max_row, cell.row)
                min_column = cell.column if min_column is None else min(min_column, cell.column)
                max_column = cell.column if max_column is None else max(max_column, cell.column)
        if None in (min_row, min_column, max_row, max_column):
            return None
        return min_row, min_column, max_row, max_column

    @staticmethod
    def _make_cell_range(min_row: int, min_column: int, max_row: int, max_column: int) -> str:
        return f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"

    @staticmethod
    def _serialize_cell_value(value: Any) -> Optional[str]:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.isoformat(sep=" ", timespec="seconds")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, time):
            return value.isoformat(timespec="seconds")
        if isinstance(value, bool):
            return "true" if value else "false"
        return normalize_text(str(value)) or None

    @staticmethod
    def _infer_title(elements: List[DocumentElement]) -> Optional[str]:
        for element in elements:
            if element.table is None:
                continue
            for row in element.table.rows:
                for cell in row:
                    if cell and len(cell) <= 150:
                        return cell.splitlines()[0]
        return None

    @staticmethod
    def _infer_document_type(relative: str) -> DocumentType:
        if "투자설명서" in relative:
            return DocumentType.INVESTMENT_PRODUCT
        if "docs_renamed" in relative:
            return DocumentType.PENSION_GUIDE
        return DocumentType.UNKNOWN
