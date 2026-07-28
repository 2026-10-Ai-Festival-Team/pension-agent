from typing import List

from openpyxl.utils.cell import get_column_letter, range_boundaries

from src.models.chunk import ChunkLocator, ChunkType, SearchChunk
from src.models.document import ParsedDocument
from src.preprocessing.common import make_chunk_id
from src.preprocessing.config import ChunkingConfig
from src.preprocessing.table_chunker import TableChunker


README_SHEET_NAMES = {"readme", "안내", "설명", "사용방법"}


class SpreadsheetChunker:
    def __init__(self, config: ChunkingConfig) -> None:
        self.config = config

    def chunk(self, document: ParsedDocument) -> List[SearchChunk]:
        chunks = []
        for element in document.elements:
            if element.table is None or not element.locator.sheet or not element.locator.cell_range:
                continue
            rows = element.table.rows
            if not rows:
                continue
            min_col, min_row, max_col, max_row = range_boundaries(element.locator.cell_range)
            if element.locator.sheet.casefold() in README_SHEET_NAMES:
                text = TableChunker.render_rows(rows)
                chunks.append(SearchChunk(chunk_id=make_chunk_id(document.source_id, ChunkType.PARAGRAPH_GROUP, [element.element_id], "sheet-description"), source_id=document.source_id, source_path=document.relative_path, source_format=document.source_format.value, document_type=document.document_type.value, title=document.title, section=element.locator.sheet, chunk_type=ChunkType.PARAGRAPH_GROUP, text=text, locator=ChunkLocator(sheet=element.locator.sheet, cell_range=element.locator.cell_range), product_codes=document.product_codes, date_candidates=document.date_candidates, element_ids=[element.element_id], metadata={"sheet_kind": "description"}))
                continue
            header_count = min(self.config.spreadsheet_header_rows, len(rows))
            header, data = rows[:header_count], rows[header_count:]
            for index in range(0, len(data), self.config.spreadsheet_rows_per_chunk):
                group = data[index:index + self.config.spreadsheet_rows_per_chunk]
                data_start = min_row + header_count + index
                data_end = data_start + len(group) - 1
                cell_range = f"{get_column_letter(min_col)}{data_start}:{get_column_letter(max_col)}{data_end}"
                header_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{min_row + header_count - 1}"
                chunks.append(SearchChunk(chunk_id=make_chunk_id(document.source_id, ChunkType.SPREADSHEET_ROWS, [element.element_id], f"rows-{data_start}-{data_end}"), source_id=document.source_id, source_path=document.relative_path, source_format=document.source_format.value, document_type=document.document_type.value, title=document.title, section=element.locator.sheet, chunk_type=ChunkType.SPREADSHEET_ROWS, text=TableChunker.render_rows(header + group), locator=ChunkLocator(sheet=element.locator.sheet, cell_range=cell_range), product_codes=document.product_codes, date_candidates=document.date_candidates, element_ids=[element.element_id], metadata={"header_range": header_range, "data_row_start": data_start, "data_row_end": data_end}))
        return chunks
