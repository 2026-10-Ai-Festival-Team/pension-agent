from pathlib import Path

from docx import Document

from src.ingestion.docx_parser import DocxParser
from src.models.document import ElementType, SourceFormat


def create_sample_docx(path: Path) -> None:
    document = Document()
    document.add_heading("퇴직연금 안내", level=0)
    document.add_paragraph("기준일: 2026.01.01")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "구분"
    table.cell(0, 1).text = "운용 주체"
    table.cell(1, 0).text = "DB"
    table.cell(1, 1).text = "회사"
    document.add_paragraph("DC는 근로자가 운용합니다.")
    document.save(path)


def test_docx_parser_preserves_block_order(tmp_path: Path) -> None:
    source_root = tmp_path / "raw"
    source_root.mkdir()
    path = source_root / "sample.docx"
    create_sample_docx(path)

    result = DocxParser().parse(path, source_root)

    assert result.source_format == SourceFormat.DOCX
    assert result.elements
    assert [element.order for element in result.elements] == list(range(len(result.elements)))
    assert [element.locator.block_index for element in result.elements] == [0, 1, 2, 3]


def test_docx_parser_extracts_table(tmp_path: Path) -> None:
    source_root = tmp_path / "raw"
    source_root.mkdir()
    path = source_root / "sample.docx"
    create_sample_docx(path)

    result = DocxParser().parse(path, source_root)
    tables = [element for element in result.elements if element.kind == ElementType.TABLE]

    assert len(tables) == 1
    assert tables[0].table is not None
    assert tables[0].table.rows[1] == ["DB", "회사"]


def test_docx_parser_does_not_invent_page_number(tmp_path: Path) -> None:
    source_root = tmp_path / "raw"
    source_root.mkdir()
    path = source_root / "sample.docx"
    create_sample_docx(path)

    result = DocxParser().parse(path, source_root)

    assert all(element.locator.page is None for element in result.elements)
    assert "2026.01.01" in result.date_candidates
