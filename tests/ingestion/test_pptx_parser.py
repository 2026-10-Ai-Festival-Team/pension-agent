from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from src.ingestion.pptx_parser import PptxParser
from src.models.document import ElementType, SourceFormat


def create_sample_pptx(path: Path) -> None:
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title_slide.shapes.title.text = "퇴직연금 상품 안내"
    text_box = title_slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    text_box.text = "기준일: 2026.01.01\n상품코드 KR5123490017"
    table = title_slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "구분"
    table.cell(0, 1).text = "위험등급"
    table.cell(1, 0).text = "채권형"
    table.cell(1, 1).text = "5등급"
    group = title_slide.shapes.add_group_shape()
    grouped_text = group.shapes.add_textbox(Inches(1), Inches(5), Inches(5), Inches(0.5))
    grouped_text.text = "그룹 내부 텍스트"
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "두 번째 슬라이드 본문"
    presentation.save(str(path))


def parse_sample(tmp_path: Path):
    source_root = tmp_path / "raw"
    source_root.mkdir()
    path = source_root / "sample.pptx"
    create_sample_pptx(path)
    return PptxParser().parse(path, source_root)


def test_pptx_parser_preserves_slide_locator(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert result.source_format == SourceFormat.PPTX
    assert {element.locator.slide for element in result.elements} == {1, 2}
    assert all(element.locator.page is None for element in result.elements)


def test_pptx_parser_extracts_title(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert result.title == "퇴직연금 상품 안내"
    assert sum(element.kind == ElementType.TITLE for element in result.elements) == 1


def test_pptx_parser_extracts_table(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    tables = [element for element in result.elements if element.kind == ElementType.TABLE]
    assert len(tables) == 1
    assert tables[0].table is not None
    assert tables[0].table.rows[1] == ["채권형", "5등급"]


def test_pptx_parser_extracts_codes_dates_and_group_text(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert "KR5123490017" in result.product_codes
    assert "2026.01.01" in result.date_candidates
    grouped = [element for element in result.elements if element.text == "그룹 내부 텍스트"]
    assert grouped and grouped[0].metadata["group_path"]


def test_pptx_parser_preserves_bbox(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert all(element.locator.bbox is not None for element in result.elements)
