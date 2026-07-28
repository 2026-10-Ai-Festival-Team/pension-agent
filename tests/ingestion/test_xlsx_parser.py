from datetime import date
from pathlib import Path

from openpyxl import Workbook

from src.ingestion.xlsx_parser import XlsxParser
from src.models.document import ElementType, SourceFormat


def create_sample_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "상품현황"
    worksheet["A1"] = "퇴직연금 상품 안내"
    worksheet["A2"] = "상품코드"
    worksheet["B2"] = "KR5123490017"
    worksheet["A3"] = "기준일"
    worksheet["B3"] = date(2026, 1, 1)
    worksheet["A4"] = "잔고"
    worksheet["B4"] = 1_000_000
    worksheet["B5"] = 500_000
    worksheet["B6"] = "=SUM(B4:B5)"
    worksheet.merge_cells("A1:B1")
    worksheet.row_dimensions[5].hidden = True
    worksheet.column_dimensions["C"].hidden = True
    hidden_sheet = workbook.create_sheet("내부자료")
    hidden_sheet.sheet_state = "hidden"
    hidden_sheet["A1"] = "숨김 시트 내용"
    workbook.save(path)
    workbook.close()


def parse_sample(tmp_path: Path):
    source_root = tmp_path / "raw"
    source_root.mkdir()
    path = source_root / "sample.xlsx"
    create_sample_xlsx(path)
    return XlsxParser().parse(path, source_root)


def test_xlsx_parser_preserves_sheet_locator(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert result.source_format == SourceFormat.XLSX
    assert len(result.elements) == 2
    assert result.elements[0].locator.sheet == "상품현황"
    assert result.elements[0].locator.cell_range == "A1:B6"
    assert all(element.locator.page is None and element.locator.slide is None for element in result.elements)


def test_xlsx_parser_extracts_merged_ranges(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert result.elements[0].kind == ElementType.TABLE
    assert result.elements[0].table is not None
    assert "A1:B1" in result.elements[0].table.merged_ranges


def test_xlsx_parser_preserves_formula(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    formula = result.elements[0].metadata["formula_cells"]["B6"]
    assert formula["formula"] == "=SUM(B4:B5)"
    assert formula["cached_value_available"] is False


def test_xlsx_parser_extracts_product_code_and_date(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert "KR5123490017" in result.product_codes
    assert "2026-01-01" in result.date_candidates


def test_xlsx_parser_records_hidden_sheet_rows_and_columns(tmp_path: Path) -> None:
    result = parse_sample(tmp_path)
    assert result.metadata["hidden_sheet_count"] == 1
    assert result.elements[1].metadata["sheet_state"] == "hidden"
    assert 5 in result.elements[0].metadata["hidden_rows"]
    assert "C" in result.elements[0].metadata["hidden_columns"]


def test_xlsx_parser_serializes_to_json(tmp_path: Path) -> None:
    assert '"source_format":"xlsx"' in parse_sample(tmp_path).model_dump_json()
